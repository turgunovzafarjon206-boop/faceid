"""HolliHop 'Расписание преподавателей' Excel faylini o'qish."""
import re
import datetime as dt
from openpyxl import load_workbook

_date_re = re.compile(r"(\d{2})\.(\d{2})\s*\(")   # 01.09 (
_time_re = re.compile(r"(\d{1,2})[:.](\d{2})")     # 8:00, 9.30, 17:00


def _teacher_name(v):
    return str(v).split("\n")[0].strip()


def parse_schedule_file(path, year=None):
    """Qaytaradi: [(sched_name, 'YYYY-MM-DD', start_min), ...] va topilgan ustozlar ro'yxati."""
    if year is None:
        year = dt.datetime.now().year
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]

    cur_date = None
    cur_teacher = None
    entries = []
    teachers = set()

    for row in ws.iter_rows(values_only=True):
        row = list(row)[:9]
        c0 = row[0]
        joined = " ".join(str(c) for c in row if c)

        # sana bloki (0-ustun bo'sh, biror katakda 'DD.MM (')
        if c0 is None:
            m = _date_re.search(joined)
            if m:
                dd, mm = m.group(1), m.group(2)
                cur_date = f"{year}-{mm}-{dd}"
                continue

        # ustoz ismi (0-ustun)
        if c0 and _teacher_name(c0) not in ("Имя", "Расписание преподавателей"):
            if not _date_re.search(str(c0)):
                cur_teacher = _teacher_name(c0)
                teachers.add(cur_teacher)

        # kataklardan vaqt
        if cur_teacher and cur_date:
            for c in row[1:]:
                if not c:
                    continue
                tm = _time_re.search(str(c))
                if tm:
                    h, mi = int(tm.group(1)), int(tm.group(2))
                    if 6 <= h <= 23 and mi < 60:
                        entries.append((cur_teacher, cur_date, h * 60 + mi))
    wb.close()
    return entries, sorted(teachers)
