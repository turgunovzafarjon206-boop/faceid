"""Kirish/chiqish tabel Excel faylini o'qish (tabel_kunlik_YYYY-MM_filial-N.xlsx).

Tuzilish: har bir varaq — bitta xodim.
  'Xodim:' qatori -> ism-familiya
  'Sana' sarlavhali jadval -> Sana (dd.mm.yyyy), Kelish vaqti, Ketish vaqti ('—' = yo'q)
"""
import re
import datetime as dt
from openpyxl import load_workbook

_EMPTY = {"", "—", "-", "–", "none"}


def _txt(v):
    return "" if v is None else str(v).strip()


def _norm_head(v):
    return _txt(v).lower().replace("ʻ", "'").replace("’", "'")


def parse_time(v):
    """'08:05' / time / datetime -> 'HH:MM' yoki None."""
    if v is None:
        return None
    if isinstance(v, dt.datetime):
        return v.strftime("%H:%M")
    if isinstance(v, dt.time):
        return v.strftime("%H:%M")
    s = _txt(v)
    if s.lower() in _EMPTY:
        return None
    m = re.match(r"^(\d{1,2})[:.](\d{2})", s)
    if not m:
        return None
    h, mi = int(m.group(1)), int(m.group(2))
    if 0 <= h <= 23 and 0 <= mi <= 59:
        return f"{h:02d}:{mi:02d}"
    return None


def parse_date(v):
    """'01.09.2026' / datetime / date -> 'YYYY-MM-DD' yoki None."""
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    s = _txt(v)
    for fmt in ("%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return dt.datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            continue
    return None


def parse_workbook(path):
    """Qaytaradi: [{'sheet', 'name', 'rows': [(day, kelish, ketish), ...]}]."""
    wb = load_workbook(path, data_only=True, read_only=True)
    result = []
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        name, header_idx, cols = None, None, {}
        for i, row in enumerate(rows[:40]):
            row = list(row or [])
            first = _norm_head(row[0] if row else "")
            if first.startswith("xodim") and len(row) > 1 and _txt(row[1]):
                name = _txt(row[1])
            if first == "sana":
                header_idx = i
                for j, h in enumerate(row):
                    h = _norm_head(h)
                    if h.startswith("kelish"):
                        cols["in"] = j
                    elif h.startswith("ketish"):
                        cols["out"] = j
                break
        if header_idx is None or "in" not in cols:
            continue
        name = name or ws.title
        data = []
        for row in rows[header_idx + 1:]:
            row = list(row or [])
            if not row:
                continue
            day = parse_date(row[0])
            if not day:
                continue
            kel = parse_time(row[cols["in"]]) if cols.get("in") is not None and len(row) > cols["in"] else None
            ket = parse_time(row[cols["out"]]) if cols.get("out") is not None and len(row) > cols["out"] else None
            if kel or ket:
                data.append((day, kel, ket))
        result.append({"sheet": ws.title, "name": name, "rows": data})
    wb.close()
    return result
